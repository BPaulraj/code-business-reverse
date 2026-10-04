#!/usr/bin/env python3
"""
Export a kit Markdown document (e.g. the Business Requirements Specification or a
review pack) to Word (.docx) and/or PDF.

    python kit/tools/export_doc.py <input.md> [--format docx|pdf|both]
                                   [--out-dir DIR] [--template corporate.docx]
                                   [--pdf-engine auto|browser|word] [--no-diagrams]

Requirements
  * Python 3.9+ and python-docx           (pip install -r kit/tools/requirements.txt)
  * PDF:  Microsoft Edge or Google Chrome (headless print)   -- or Microsoft Word
  * Diagrams: rendered offline by the browser with kit/tools/vendor/mermaid.min.js.
    Without a browser, diagram source is shown as a code block instead.
Nothing is sent over the network.
"""
from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
MERMAID_JS = HERE / "vendor" / "mermaid.min.js"

# --------------------------------------------------------------------------- markdown parsing

LIST_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


def split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in cells]


def parse_markdown(text: str) -> list[tuple]:
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    if text.startswith("---\n"):
        end = text.find("\n---", 4)
        if end != -1:
            text = text[end + 4:]
    lines = text.splitlines()
    blocks: list[tuple] = []
    i, n = 0, len(lines)
    para: list[str] = []

    def flush_para():
        if para:
            blocks.append(("para", " ".join(s.strip() for s in para)))
            para.clear()

    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_para(); i += 1; continue
        if stripped.startswith("```"):
            flush_para()
            lang = stripped[3:].strip().lower()
            body = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                body.append(lines[i]); i += 1
            i += 1
            blocks.append(("mermaid" if lang == "mermaid" else "code", lang, "\n".join(body)))
            continue
        m = HEADING_RE.match(line)
        if m:
            flush_para()
            blocks.append(("heading", len(m.group(1)), m.group(2)))
            i += 1; continue
        if re.match(r"^\s*(---+|\*\*\*+|___+)\s*$", line):
            flush_para(); blocks.append(("hr",)); i += 1; continue
        if stripped.startswith("|") and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]):
            flush_para()
            header = split_row(line)
            rows = []
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            blocks.append(("table", header, rows))
            continue
        if stripped.startswith(">"):
            flush_para()
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip()); i += 1
            blocks.append(("quote", " ".join(q for q in quote if q)))
            continue
        lm = LIST_RE.match(line)
        if lm:
            flush_para()
            items = []
            while i < n:
                lm = LIST_RE.match(lines[i])
                if lm:
                    indent = len(lm.group(1).expandtabs(4))
                    level = 0 if indent < 2 else (1 if indent < 5 else 2)
                    ordered = lm.group(2)[0].isdigit()
                    number = int(re.match(r"\d+", lm.group(2)).group()) if ordered else None
                    items.append([level, ordered, number, lm.group(3).strip()])
                    i += 1
                elif lines[i].strip() and lines[i].startswith((" ", "\t")) and items \
                        and not lines[i].strip().startswith(("|", "```")):
                    items[-1][3] += " " + lines[i].strip(); i += 1
                else:
                    break
            blocks.append(("list", items))
            continue
        para.append(line)
        i += 1
    flush_para()
    return blocks


INLINE_RE = re.compile(
    r"(`[^`]+`)"                          # code
    r"|(\[[^\]]+\]\([^)]+\))"             # link
    r"|(\*\*.+?\*\*)"                     # bold
    r"|((?<![\w*])\*(?!\s)[^*]+?\*(?![\w*]))"  # italic
)


def inline_tokens(text: str, bold=False, italic=False) -> list[dict]:
    """Return runs: {text, bold, italic, code, link}."""
    out, pos = [], 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append(dict(text=text[pos:m.start()], bold=bold, italic=italic, code=False, link=None))
        tok = m.group(0)
        if m.group(1):
            out.append(dict(text=tok[1:-1], bold=bold, italic=italic, code=True, link=None))
        elif m.group(2):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok).groups()
            for r in inline_tokens(label, bold, italic):
                r["link"] = url
                out.append(r)
        elif m.group(3):
            out.extend(inline_tokens(tok[2:-2], True, italic))
        else:
            out.extend(inline_tokens(tok[1:-1], bold, True))
        pos = m.end()
    if pos < len(text):
        out.append(dict(text=text[pos:], bold=bold, italic=italic, code=False, link=None))
    for r in out:
        r["text"] = html.unescape(r["text"]).replace("<br/>", "\n").replace("<br>", "\n")
    return out


# --------------------------------------------------------------------------- browser helpers

def find_browser() -> str | None:
    candidates = [
        os.environ.get("KIT_BROWSER", ""),
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    ]
    for name in ("msedge", "microsoft-edge", "google-chrome", "chrome", "chromium", "chromium-browser"):
        p = shutil.which(name)
        if p:
            candidates.append(p)
    probe_dir = Path(tempfile.mkdtemp(prefix="kit-probe-"))
    probe = probe_dir / "probe.html"
    probe.write_text('<p id="x">no</p><script>document.getElementById("x").textContent="ok-js"</script>',
                     encoding="utf-8")
    try:
        seen = set()
        for c in candidates:
            if not c or c in seen or not Path(c).exists():
                continue
            seen.add(c)
            # Some installs (policy-restricted or mid-update) start but render nothing: test first.
            try:
                r = run_browser(c, ["--dump-dom", probe.as_uri()], timeout=60)
                if "ok-js" in (r.stdout or ""):
                    return c
                print(f"  ! {Path(c).name} did not respond in headless mode; trying next browser", file=sys.stderr)
            except (OSError, subprocess.TimeoutExpired):
                continue
    finally:
        shutil.rmtree(probe_dir, ignore_errors=True)
    return None


def run_browser(browser: str, args: list[str], timeout=180) -> subprocess.CompletedProcess:
    profile = tempfile.mkdtemp(prefix="kit-browser-")
    try:
        cmd = [browser, "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
               "--disable-extensions", "--allow-file-access-from-files", f"--user-data-dir={profile}", *args]
        return subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace",
                              timeout=timeout)
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def render_mermaid(diagrams: list[str], browser: str | None, workdir: Path) -> list[dict | None]:
    """Render each Mermaid source to {svg, width, height, png}. None where rendering failed."""
    if not diagrams or not browser or not MERMAID_JS.exists():
        return [None] * len(diagrams)
    page = workdir / "mermaid-render.html"
    page.write_text(f"""<!doctype html><html><head><meta charset="utf-8">
<script src="{MERMAID_JS.as_uri()}"></script>
<style>body{{margin:0;background:#fff}} .d{{display:inline-block;padding:8px}}</style></head><body>
<pre id="out">pending</pre>
<script>
const SRC = {json.dumps(diagrams)};
(async () => {{
  const res = [];
  try {{
    mermaid.initialize({{startOnLoad:false, theme:'neutral', securityLevel:'loose',
                        flowchart:{{htmlLabels:false}}, fontFamily:'Segoe UI, Arial, sans-serif'}});
    for (let i = 0; i < SRC.length; i++) {{
      try {{
        const {{svg}} = await mermaid.render('m' + i, SRC[i]);
        const div = document.createElement('div'); div.className = 'd'; div.innerHTML = svg;
        document.body.appendChild(div);
        const el = div.querySelector('svg'); const r = el.getBoundingClientRect();
        res.push({{ok:true, svg:el.outerHTML, width:Math.ceil(r.width), height:Math.ceil(r.height)}});
      }} catch (e) {{ res.push({{ok:false, error:String(e)}}); }}
    }}
  }} catch (e) {{ res.push({{fatal:String(e)}}); }}
  document.getElementById('out').textContent = btoa(unescape(encodeURIComponent(JSON.stringify(res))));
}})();
</script></body></html>""", encoding="utf-8")
    proc = run_browser(browser, ["--virtual-time-budget=30000", "--dump-dom", page.as_uri()])
    m = re.search(r'<pre id="out">([A-Za-z0-9+/=]+)</pre>', proc.stdout or "")
    if not m:
        print("  ! diagram rendering failed; diagram source will be shown instead", file=sys.stderr)
        return [None] * len(diagrams)
    import base64
    results = json.loads(base64.b64decode(m.group(1)).decode("utf-8"))
    out: list[dict | None] = []
    for i, r in enumerate(results[: len(diagrams)]):
        if not r.get("ok"):
            print(f"  ! diagram {i + 1} not rendered: {r.get('error') or r.get('fatal')}", file=sys.stderr)
            out.append(None)
            continue
        w, h = max(r["width"], 50), max(r["height"], 30)
        # Screenshot each diagram at 2x for a crisp PNG (used in Word).
        shot_html = workdir / f"diagram-{i}.html"
        shot_html.write_text(
            f'<!doctype html><html><head><meta charset="utf-8"><style>html,body{{margin:0;padding:0;'
            f'background:#fff;overflow:hidden}} svg{{display:block}}</style></head><body>{r["svg"]}</body></html>',
            encoding="utf-8")
        png = workdir / f"diagram-{i}.png"
        run_browser(browser, ["--force-device-scale-factor=2", "--hide-scrollbars",
                              f"--window-size={w},{h}", f"--screenshot={png}", shot_html.as_uri()])
        out.append(dict(svg=r["svg"], width=w, height=h, png=png if png.exists() else None))
    while len(out) < len(diagrams):
        out.append(None)
    return out


# --------------------------------------------------------------------------- HTML (for PDF)

CSS = """
@page { size: A4; margin: 18mm 16mm 20mm 16mm;
        @bottom-center { content: "Page " counter(page) " of " counter(pages); font: 8pt 'Segoe UI', Arial; color: #666; }
        @top-right { content: string(doctitle); font: 8pt 'Segoe UI', Arial; color: #666; } }
@page :first { @top-right { content: none; } @bottom-center { content: none; } }
body { font: 10pt/1.45 'Segoe UI', Calibri, Arial, sans-serif; color: #1f2328; }
.cover { height: 240mm; display: flex; flex-direction: column; justify-content: center; break-after: page; }
.cover h1 { font-size: 26pt; margin: 0 0 8mm; color: #1f3a68; string-set: doctitle content(); }
.cover .meta { color: #555; font-size: 10pt; }
.toc { break-after: page; } .toc h2 { border: none; } .toc ol { padding-left: 1.2em; } .toc li { margin: 2px 0; }
.toc a { color: #1f2328; text-decoration: none; }
h1, h2, h3, h4 { color: #1f3a68; break-after: avoid; }
h2 { font-size: 15pt; border-bottom: 1px solid #c9d1d9; padding-bottom: 3px; margin-top: 20px; }
h3 { font-size: 12pt; margin-top: 16px; } h4 { font-size: 10.5pt; }
table { border-collapse: collapse; width: 100%; margin: 6px 0 12px; font-size: 8.5pt; }
th, td { border: 1px solid #c9d1d9; padding: 4px 6px; vertical-align: top; text-align: left; }
th { background: #e8eef7; } tr { break-inside: avoid; } thead { display: table-header-group; }
code { font-family: Consolas, 'Courier New', monospace; font-size: 8.5pt; background: #f3f4f6; padding: 0 2px; border-radius: 2px; }
pre { background: #f6f8fa; border: 1px solid #e1e4e8; padding: 8px; font-size: 8pt; white-space: pre-wrap; break-inside: avoid; }
blockquote { margin: 8px 0; padding: 8px 12px; background: #fff8e5; border-left: 4px solid #d4a72c; }
.diagram { text-align: center; margin: 10px 0; break-inside: avoid; } .diagram svg { max-width: 100%; height: auto; }
hr { border: none; border-top: 1px solid #e1e4e8; margin: 14px 0; }
ul, ol { margin: 4px 0 8px; padding-left: 22px; } li { margin: 1px 0; }
a { color: #0b5cad; }
"""


def inline_html(text: str) -> str:
    parts = []
    for r in inline_tokens(text):
        t = html.escape(r["text"]).replace("\n", "<br>")
        if r["code"]:
            t = f"<code>{t}</code>"
        if r["bold"]:
            t = f"<strong>{t}</strong>"
        if r["italic"]:
            t = f"<em>{t}</em>"
        if r["link"] and r["link"].startswith(("http://", "https://")):
            t = f'<a href="{html.escape(r["link"])}">{t}</a>'
        parts.append(t)
    return "".join(parts)


def slug(text: str, used: set) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "section"
    base, k = s, 2
    while s in used:
        s = f"{base}-{k}"; k += 1
    used.add(s)
    return s


def build_html(blocks, diagrams, title, source_name) -> str:
    body, toc, used, d_idx, title_done = [], [], set(), 0, False
    for b in blocks:
        kind = b[0]
        if kind == "heading":
            level, text = b[1], b[2]
            if level == 1 and not title_done:
                title_done = True
                continue  # rendered on cover
            anchor = slug(text, used)
            if level in (2, 3):
                toc.append((level, text, anchor))
            body.append(f'<h{level} id="{anchor}">{inline_html(text)}</h{level}>')
        elif kind == "para":
            body.append(f"<p>{inline_html(b[1])}</p>")
        elif kind == "quote":
            body.append(f"<blockquote>{inline_html(b[1])}</blockquote>")
        elif kind == "hr":
            body.append("<hr>")
        elif kind == "code":
            body.append(f"<pre>{html.escape(b[2])}</pre>")
        elif kind == "mermaid":
            d = diagrams[d_idx]; d_idx += 1
            body.append(f'<div class="diagram">{d["svg"]}</div>' if d else
                        f'<pre>Diagram (Mermaid source):\n{html.escape(b[2])}</pre>')
        elif kind == "table":
            head = "".join(f"<th>{inline_html(c)}</th>" for c in b[1])
            rows = "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in b[2])
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
        elif kind == "list":
            body.append(list_html(b[1]))
    toc_html = "".join(
        f'<li style="margin-left:{(lvl - 2) * 16}px"><a href="#{a}">{inline_html(t)}</a></li>' for lvl, t, a in toc)
    cover = (f'<section class="cover"><h1>{inline_html(title)}</h1><div class="meta">'
             f'Generated {dt.date.today().isoformat()} from <code>{html.escape(source_name)}</code></div></section>')
    toc_sec = f'<section class="toc"><h2>Contents</h2><ol style="list-style:none;padding:0">{toc_html}</ol></section>' if toc else ""
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>'
            f"<style>{CSS}</style></head><body>{cover}{toc_sec}{''.join(body)}</body></html>")


def list_html(items) -> str:
    out, stack = [], []  # stack of (level, tag)
    for level, ordered, number, text in items:
        tag = "ol" if ordered else "ul"
        while stack and stack[-1][0] > level:
            out.append(f"</li></{stack.pop()[1]}>")
        if not stack or stack[-1][0] < level:
            start = f' start="{number}"' if ordered and number and number != 1 else ""
            out.append(f"<{tag}{start}><li>")
            stack.append((level, tag))
        else:
            out.append("</li><li>")
        out.append(inline_html(text))
    while stack:
        out.append(f"</li></{stack.pop()[1]}>")
    return "".join(out)


# --------------------------------------------------------------------------- DOCX

def build_docx(blocks, diagrams, title, source_name, out_path: Path, template: Path | None):
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for template users)
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document(str(template)) if template else Document()
    style_names = {s.name for s in doc.styles}

    def style(name, fallback="Normal"):
        return name if name in style_names else fallback

    if not template:
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        for sec in doc.sections:  # A4, sensible margins
            sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
            sec.left_margin = sec.right_margin = Inches(0.8)
            sec.top_margin = sec.bottom_margin = Inches(0.8)

    def add_field(paragraph, instr):
        run = paragraph.add_run()
        for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate"),
                          ("w:t", None), ("w:fldChar", "end")):
            el = OxmlElement(tag)
            if tag == "w:fldChar":
                el.set(qn("w:fldCharType"), attr)
            elif tag == "w:instrText":
                el.set(qn("xml:space"), "preserve"); el.text = instr
            else:
                el.text = "1" if "PAGE" in instr else "Right-click → Update field to build the table of contents."
            run._r.append(el)

    def add_runs(paragraph, text, size=None, bold=False):
        for r in inline_tokens(text, bold=bold):
            run = paragraph.add_run(r["text"])
            run.bold = r["bold"] or None
            run.italic = r["italic"] or None
            if r["code"]:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif size:
                run.font.size = size
            if r["link"]:
                run.font.color.rgb = RGBColor(0x0B, 0x5C, 0xAD)
                run.underline = True

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    # Cover page
    p = doc.add_paragraph(style=style("Title"))
    add_runs(p, title)
    meta = doc.add_paragraph()
    meta.add_run(f"Generated {dt.date.today().isoformat()} from {source_name}").italic = True
    meta.add_run().add_break(WD_BREAK.PAGE)
    # Table of contents (Word builds it on open / F9)
    doc.add_paragraph("Contents", style=style("TOC Heading", style("Heading 1")))
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Header / footer
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]; hp.text = title; hp.alignment = 2
    for r in hp.runs:
        r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    fp = sec.footer.paragraphs[0]; fp.alignment = 1
    fp.add_run("Page ").font.size = Pt(8); add_field(fp, "PAGE")
    fp.add_run(" of ").font.size = Pt(8); add_field(fp, "NUMPAGES")

    d_idx, title_done = 0, False
    for b in blocks:
        kind = b[0]
        if kind == "heading":
            level, text = b[1], b[2]
            if level == 1 and not title_done:
                title_done = True; continue
            h = doc.add_paragraph(style=style(f"Heading {max(1, min(level - 1, 4))}"))
            add_runs(h, text)
        elif kind == "para":
            add_runs(doc.add_paragraph(), b[1])
        elif kind == "quote":
            q = doc.add_paragraph(style=style("Intense Quote", style("Quote")))
            add_runs(q, b[1])
        elif kind == "hr":
            continue
        elif kind in ("code", "mermaid"):
            d = diagrams[d_idx] if kind == "mermaid" else None
            if kind == "mermaid":
                d_idx += 1
            if d and d.get("png"):
                width = min(6.4, d["width"] / 96)
                doc.add_picture(str(d["png"]), width=Inches(width))
                doc.paragraphs[-1].alignment = 1
            else:
                cp = doc.add_paragraph()
                run = cp.add_run(("Diagram (Mermaid source):\n" if kind == "mermaid" else "") + b[2])
                run.font.name = "Consolas"; run.font.size = Pt(8.5)
        elif kind == "table":
            header, rows = b[1], b[2]
            ncols = max(len(header), *(len(r) for r in rows)) if rows else len(header)
            t = doc.add_table(rows=1, cols=ncols)
            t.style = style("Table Grid", t.style.name if t.style else "Normal")
            for j in range(ncols):
                cell = t.rows[0].cells[j]
                cell.text = ""
                add_runs(cell.paragraphs[0], header[j] if j < len(header) else "", size=Pt(8.5), bold=True)
                shade(cell, "E8EEF7")
            trPr = t.rows[0]._tr.get_or_add_trPr()  # repeat header row on each page
            hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trPr.append(hdr)
            for r in rows:
                cells = t.add_row().cells
                for j in range(ncols):
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0], r[j] if j < len(r) else "", size=Pt(8.5))
            doc.add_paragraph()
        elif kind == "list":
            counters: dict[int, int] = {}
            for level, ordered, number, text in b[1]:
                if ordered:
                    counters[level] = (number or counters.get(level, 0) + 1)
                    p = doc.add_paragraph(style=style("List Paragraph"))
                    p.paragraph_format.left_indent = Inches(0.3 + 0.3 * level)
                    p.paragraph_format.first_line_indent = Inches(-0.22)
                    p.add_run(f"{counters[level]}. ")
                else:
                    name = "List Bullet" if level == 0 else f"List Bullet {level + 1}"
                    p = doc.add_paragraph(style=style(name, style("List Bullet", "List Paragraph")))
                add_runs(p, text)

    # Ask Word to refresh TOC / page fields when the document is opened.
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); settings.append(upd)
    doc.core_properties.title = title
    doc.save(str(out_path))


# --------------------------------------------------------------------------- Word automation (optional)

def word_available() -> bool:
    if os.name != "nt":
        return False
    try:
        r = subprocess.run(["reg", "query", r"HKCR\Word.Application"], capture_output=True, text=True)
        return r.returncode == 0
    except OSError:
        return False


def word_finalize(docx_path: Path, pdf_path: Path | None) -> bool:
    """Open the .docx in Word, update TOC and fields, save, and optionally export PDF."""
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false; $w.DisplayAlerts = 0
try {{
  $d = $w.Documents.Open('{docx_path}', $false, $false)
  foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
  $d.Fields.Update() | Out-Null
  $d.Save()
  {"$d.SaveAs2('" + str(pdf_path) + "', 17)" if pdf_path else ""}
  $d.Close($false)
}} finally {{ $w.Quit() }}
"""
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                       capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        print("  ! Word automation failed: " + (r.stderr or r.stdout).strip()[:400], file=sys.stderr)
    return r.returncode == 0


# --------------------------------------------------------------------------- main

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Export a kit Markdown document to DOCX and/or PDF.")
    ap.add_argument("input", type=Path)
    ap.add_argument("--format", choices=["docx", "pdf", "both"], default="both")
    ap.add_argument("--out-dir", type=Path)
    ap.add_argument("--template", type=Path, help="corporate .docx whose styles are reused")
    ap.add_argument("--pdf-engine", choices=["auto", "browser", "word"], default="auto")
    ap.add_argument("--no-diagrams", action="store_true", help="show Mermaid source instead of rendering")
    ap.add_argument("--no-word", action="store_true", help="never automate Microsoft Word")
    ap.add_argument("--keep-html", action="store_true", help="also write the print-ready .html next to the outputs")
    a = ap.parse_args(argv)

    src = a.input.resolve()
    if not src.exists():
        print(f"Input not found: {src}", file=sys.stderr); return 2
    out_dir = (a.out_dir or src.parent).resolve(); out_dir.mkdir(parents=True, exist_ok=True)
    blocks = parse_markdown(src.read_text(encoding="utf-8"))
    title = next((b[2] for b in blocks if b[0] == "heading" and b[1] == 1), src.stem)
    title_plain = "".join(r["text"] for r in inline_tokens(title))

    want_docx = a.format in ("docx", "both")
    want_pdf = a.format in ("pdf", "both")
    browser = find_browser()
    word = word_available() and not a.no_word
    engine = a.pdf_engine
    if want_pdf and engine == "auto":
        engine = "browser" if browser else ("word" if word else None)
    if want_pdf and engine is None:
        print("No PDF engine: install/point KIT_BROWSER at Edge or Chrome, or install Microsoft Word.", file=sys.stderr)
        return 3
    if want_pdf and engine == "browser" and not browser:
        print("Browser engine requested but Edge/Chrome not found (set KIT_BROWSER).", file=sys.stderr); return 3
    if want_pdf and engine == "word" and not word:
        print("Word engine requested but Microsoft Word is not available.", file=sys.stderr); return 3

    mermaid_src = [b[2] for b in blocks if b[0] == "mermaid"]
    work = Path(tempfile.mkdtemp(prefix="kit-export-"))
    try:
        diagrams = [None] * len(mermaid_src) if a.no_diagrams else render_mermaid(mermaid_src, browser, work)
        rendered = sum(1 for d in diagrams if d)
        if mermaid_src:
            print(f"Diagrams: {rendered}/{len(mermaid_src)} rendered")

        written = []
        docx_path = out_dir / f"{src.stem}.docx"
        pdf_path = out_dir / f"{src.stem}.pdf"
        if want_docx or (want_pdf and engine == "word"):
            build_docx(blocks, diagrams, title_plain, src.name, docx_path, a.template)
            if word and (want_docx or engine == "word"):
                ok = word_finalize(docx_path, pdf_path if (want_pdf and engine == "word") else None)
                if ok:
                    print("Word: table of contents and page numbers updated")
            if want_docx:
                written.append(docx_path)
            if want_pdf and engine == "word" and pdf_path.exists():
                written.append(pdf_path)
            if not want_docx and docx_path.exists():
                docx_path.unlink()
        if want_pdf and engine == "browser":
            html_path = work / f"{src.stem}.html"
            html_path.write_text(build_html(blocks, diagrams, title_plain, src.name), encoding="utf-8")
            if pdf_path.exists():
                pdf_path.unlink()
            run_browser(browser, ["--no-pdf-header-footer", "--print-to-pdf-no-header",
                                  f"--print-to-pdf={pdf_path}", html_path.as_uri()])
            if pdf_path.exists():
                written.append(pdf_path)
            else:
                print("  ! browser did not produce a PDF", file=sys.stderr)
        if a.keep_html:
            keep = out_dir / f"{src.stem}.html"
            keep.write_text(build_html(blocks, diagrams, title_plain, src.name), encoding="utf-8")
            written.append(keep)
        for w in written:
            print(f"Wrote {w} ({w.stat().st_size // 1024} KB)")
        return 0 if written else 1
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
